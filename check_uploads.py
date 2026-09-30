"""
check_uploads.py — Check upload status and fix missing uploads

Commands:
    python check_uploads.py          — show full status report
    python check_uploads.py --fix    — re-upload anything missing
"""

import os
import sys
import json
from datetime import datetime
from openpyxl import load_workbook, Workbook
from config import LOG_FILE, OUTPUT_DIR

def load_tracker() -> list:
    """Load all rows from Excel tracker."""
    if not os.path.exists(LOG_FILE):
        print("  No tracker found at", LOG_FILE)
        return []
    rows = []
    try:
        wb = load_workbook(LOG_FILE)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[9]:  # no video file = empty row
                continue
            rows.append({
                "row":        row[0] if row[0] else "?",
                "date":       str(row[0] or ""),
                "time":       str(row[1] or ""),
                "video_num":  row[2],
                "topic":      str(row[3] or ""),
                "title":      str(row[4] or ""),
                "hashtags":   str(row[6] or "#shorts #viral #facts"),
                "yt_url":     str(row[7] or "").strip(),
                "ig_url":     str(row[8] or "").strip(),
                "video_file": str(row[9] or "").strip(),
                "status":     str(row[10] or "").strip(),
            })
    except Exception as e:
        print(f"  Error reading Excel: {e}")
    return rows

def check_status(rows: list) -> dict:
    """Categorize each video by upload status."""
    result = {
        "fully_uploaded":    [],
        "youtube_only":      [],
        "instagram_only":    [],
        "not_uploaded":      [],
        "file_missing":      [],
    }

    for r in rows:
        vf = r["video_file"]
        yt = r["yt_url"] not in ("", "pending", "None")
        ig = r["ig_url"] not in ("", "pending", "None")
        exists = os.path.exists(vf)

        if not exists:
            result["file_missing"].append(r)
        elif yt and ig:
            result["fully_uploaded"].append(r)
        elif yt and not ig:
            result["youtube_only"].append(r)
        elif ig and not yt:
            result["instagram_only"].append(r)
        else:
            result["not_uploaded"].append(r)

    return result

def print_report(rows: list, status: dict):
    print("\n" + "="*55)
    print("  UPLOAD STATUS REPORT")
    print(f"  {datetime.now().strftime('%A %d %B %Y, %H:%M')}")
    print("="*55)
    print(f"\n  Total videos tracked : {len(rows)}")
    print(f"  Fully uploaded       : {len(status['fully_uploaded'])} ✅")
    print(f"  YouTube only         : {len(status['youtube_only'])} ⚠️  (Instagram missing)")
    print(f"  Instagram only       : {len(status['instagram_only'])} ⚠️  (YouTube missing)")
    print(f"  Not uploaded at all  : {len(status['not_uploaded'])} ❌")
    print(f"  File missing on disk : {len(status['file_missing'])} 🗑️")

    if status["youtube_only"]:
        print("\n  ── YouTube only (need Instagram upload) ──")
        for r in status["youtube_only"]:
            print(f"    • {r['title'][:40]}")
            print(f"      YT : {r['yt_url'][:60]}")
            print(f"      IG : NOT UPLOADED")
            print(f"      File: {os.path.basename(r['video_file'])}")

    if status["instagram_only"]:
        print("\n  ── Instagram only (need YouTube upload) ──")
        for r in status["instagram_only"]:
            print(f"    • {r['title'][:40]}")
            print(f"      YT : NOT UPLOADED")
            print(f"      IG : {r['ig_url'][:60]}")
            print(f"      File: {os.path.basename(r['video_file'])}")

    if status["not_uploaded"]:
        print("\n  ── Not uploaded anywhere ──")
        for r in status["not_uploaded"]:
            print(f"    • {r['title'][:40]}")
            print(f"      File: {os.path.basename(r['video_file'])}")

    if status["file_missing"]:
        print("\n  ── File missing on disk (cannot re-upload) ──")
        for r in status["file_missing"]:
            print(f"    • {r['title'][:40]}")
            print(f"      Was at: {r['video_file']}")

    print()

def fix_missing(status: dict):
    """Re-upload anything that is missing YouTube or Instagram."""
    to_fix = (
        status["youtube_only"] +
        status["instagram_only"] +
        status["not_uploaded"]
    )

    if not to_fix:
        print("  Nothing to fix — all videos fully uploaded!")
        return

    print(f"\n  Found {len(to_fix)} video(s) with missing uploads.")
    print("  Building fix queue...\n")

    from upload_youtube import upload_to_youtube
    from upload_instagram import upload_to_instagram
    from excel_logger import update_urls

    for r in to_fix:
        vf = r["video_file"]
        title = r["title"]
        hashtags = r["hashtags"]

        if not os.path.exists(vf):
            print(f"  SKIP — file not found: {os.path.basename(vf)}")
            continue

        print(f"\n  Processing: {title[:40]}")
        print(f"  File: {os.path.basename(vf)}")

        yt_url = r["yt_url"] if r["yt_url"] not in ("", "pending", "None") else ""
        ig_url = r["ig_url"] if r["ig_url"] not in ("", "pending", "None") else ""

        # upload to YouTube if missing
        if not yt_url:
            print("  Uploading to YouTube...")
            try:
                description = f"{title}\n\n{hashtags}\n\n#shorts #youtubeshorts"
                yt_url = upload_to_youtube(vf, title, description, hashtags)
                print(f"  ✅ YouTube: {yt_url}")
            except Exception as e:
                print(f"  ❌ YouTube failed: {e}")

        # upload to Instagram if missing
        if not ig_url:
            print("  Uploading to Instagram...")
            try:
                ig_url = upload_to_instagram(vf, title, hashtags)
                print(f"  ✅ Instagram: {ig_url}")
            except Exception as e:
                print(f"  ❌ Instagram failed: {e}")

        # update Excel
        if yt_url or ig_url:
            update_urls(vf, yt_url=yt_url or None, ig_url=ig_url or None)

    print("\n  Fix complete! Check tracker.xlsx for updated URLs.")

def main():
    auto_fix = "--fix" in sys.argv

    rows = load_tracker()
    if not rows:
        print("  No videos tracked yet.")
        return

    status = check_status(rows)
    print_report(rows, status)

    needs_fix = (
        len(status["youtube_only"]) +
        len(status["instagram_only"]) +
        len(status["not_uploaded"])
    )

    if needs_fix > 0:
        if auto_fix:
            fix_missing(status)
        else:
            print(f"  {needs_fix} video(s) need attention.")
            print("  Run with --fix to automatically re-upload missing ones:")
            print("  python check_uploads.py --fix")
    else:
        print("  All good! Every video is fully uploaded.")

if __name__ == "__main__":
    main()
