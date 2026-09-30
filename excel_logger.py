import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from config import LOG_FILE

HEADERS = [
    "Date", "Time", "Video #", "Topic", "Title",
    "Script (preview)", "Hashtags",
    "YouTube URL", "Instagram URL",
    "Video File", "Status"
]

HEADER_COLOR = "1A1A2E"
ROW_A_COLOR  = "F8F9FA"
ROW_B_COLOR  = "FFFFFF"
ACCENT_COLOR = "E63946"

def _thin_border():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)

def _init_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Video Tracker"

    # header row
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", start_color=HEADER_COLOR)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()
    ws.row_dimensions[1].height = 32

    # column widths
    widths = [12, 8, 8, 20, 40, 45, 35, 40, 40, 30, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    return wb, ws

def log_video(entry: dict):
    """
    Append a video entry to the Excel tracker.
    entry keys: topic, title, script, hashtags, yt_url, ig_url, video_file, status
    """
    os.makedirs(os.path.dirname(LOG_FILE), exist_ok=True)

    if os.path.exists(LOG_FILE):
        wb = load_workbook(LOG_FILE)
        ws = wb.active
        next_row = ws.max_row + 1
        video_num = next_row - 1
    else:
        wb, ws = _init_workbook()
        next_row = 2
        video_num = 1

    now = datetime.now()
    row_color = ROW_A_COLOR if video_num % 2 == 0 else ROW_B_COLOR

    values = [
        now.strftime("%Y-%m-%d"),
        now.strftime("%H:%M"),
        video_num,
        entry.get("topic", ""),
        entry.get("title", ""),
        entry.get("script", "")[:120] + "…",
        entry.get("hashtags", ""),
        entry.get("yt_url", "pending"),
        entry.get("ig_url", "pending"),
        entry.get("video_file", ""),
        entry.get("status", "created"),
    ]

    for col, val in enumerate(values, 1):
        cell = ws.cell(row=next_row, column=col, value=val)
        cell.fill = PatternFill("solid", start_color=row_color)
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _thin_border()

    # hyperlinks
    if entry.get("yt_url") and entry["yt_url"] != "pending":
        ws.cell(row=next_row, column=8).hyperlink = entry["yt_url"]
        ws.cell(row=next_row, column=8).font = Font(name="Arial", size=10, color="0563C1", underline="single")
    if entry.get("ig_url") and entry["ig_url"] != "pending":
        ws.cell(row=next_row, column=9).hyperlink = entry["ig_url"]
        ws.cell(row=next_row, column=9).font = Font(name="Arial", size=10, color="0563C1", underline="single")

    ws.row_dimensions[next_row].height = 28
    wb.save(LOG_FILE)
    print(f"  Logged to Excel → row {next_row}")


def update_urls(video_file: str, yt_url: str = None, ig_url: str = None):
    """After upload, update the URL columns for a given video file."""
    if not os.path.exists(LOG_FILE):
        return
    wb = load_workbook(LOG_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[9].value == video_file:   # column J = video file
            if yt_url:
                row[7].value = yt_url    # column H
                row[7].hyperlink = yt_url
                row[7].font = Font(name="Arial", size=10, color="0563C1", underline="single")
            if ig_url:
                row[8].value = ig_url    # column I
                row[8].hyperlink = ig_url
                row[8].font = Font(name="Arial", size=10, color="0563C1", underline="single")
            row[10].value = "uploaded"   # column K = status
            break
    wb.save(LOG_FILE)
    print(f"  Excel URLs updated.")
